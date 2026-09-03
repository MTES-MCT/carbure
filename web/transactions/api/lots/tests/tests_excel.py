import os
from datetime import date
from unittest.mock import patch

import openpyxl
import pandas as pd
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from core.models import CarbureLot, Entity, Pays
from core.serializers import CarbureLotCSVSerializer
from core.tests_utils import setup_current_user
from core.xlsx_v3 import template_v4
from transactions.helpers import INVALID_DISPATCH_SITE, fill_dispatch_data
from transactions.models import Depot, ProductionSite
from transactions.models.entity_site import EntitySite
from transactions.models.site import Site


class LotsExcelImportTest(TestCase):
    fixtures = [
        "json/biofuels.json",
        "json/feedstock.json",
        "json/countries.json",
    ]

    # create the entity that will import the excel file
    def setupOwnerEntity(self):
        self.owner = Entity.objects.create(
            is_enabled=True,
            name="Owner",
            default_certificate="OWNER_CERTIFICATE",
        )
        self.owner_production_site = ProductionSite.objects.create(
            site_type=Site.PRODUCTION_BIOLIQUID,
            is_enabled=True,
            name="Owner Production Site",
            country=self.FR,
            created_by=self.owner,
        )
        self.owner_depot = Depot.objects.create(
            site_type=Site.EFPE,
            is_enabled=True,
            name="Owner Depot",
            customs_id="AAA",
            country=self.FR,
        )

    # create another entity that will be involved in transactions
    def setupOtherEntity(self):
        self.other = Entity.objects.create(
            is_enabled=True,
            name="Other",
            entity_type=Entity.OPERATOR,
            default_certificate="OTHER_CERTIFICATE",
            has_trading=True,
        )
        self.other_depot = Depot.objects.create(
            site_type=Site.EFPE,
            is_enabled=True,
            name="Other Depot",
            customs_id="BBB",
            country=self.FR,
        )
        EntitySite.objects.create(entity=self.other, site=self.other_depot)

    def setUp(self):
        self.FR = Pays.objects.get(code_pays="FR")

        self.setupOwnerEntity()
        self.setupOtherEntity()

        self.user = setup_current_user(
            self,
            "tester@carbure.local",
            "Tester",
            "gogogo",
            [(self.owner, "ADMIN")],
        )

    def debug_lots(self, title, lots):
        values = lots.values(
            "free_field",
            "carbure_producer",
            "unknown_producer",
            "carbure_production_site",
            "unknown_production_site",
            "carbure_supplier",
            "unknown_supplier",
            "carbure_vendor",
            "carbure_client",
            "unknown_client",
        )

        print(title)
        print(pd.DataFrame(values).fillna(""))

    def send_excel(self, entity: Entity, excel_fixture: str):
        CarbureLot.objects.filter(added_by=entity).delete()

        filepath = f"{os.environ['CARBURE_HOME']}/web/transactions/fixtures/{excel_fixture}"
        with open(filepath, "rb") as reader:
            file = SimpleUploadedFile(excel_fixture, reader.read())

        response = self.client.post(
            reverse("transactions-lots-add-excel"),
            {"entity_id": entity.id, "file": file},
        )
        if response.status_code != 200:
            print(response.json(), response.status_code)

        return CarbureLot.objects.filter(added_by=entity).all().prefetch_related("genericerror_set")

    def test_producer_excel(self):
        self.owner.entity_type = Entity.PRODUCER
        self.owner.has_trading = False
        self.owner.has_stocks = False
        self.owner.save()

        # add production site, remove depot
        EntitySite.objects.update_or_create(entity=self.owner, site=self.owner_production_site)
        EntitySite.objects.filter(entity=self.owner, site=self.owner_depot).delete()

        lots = self.send_excel(self.owner, "test_lot_template.xlsx")

        assert lots.filter(carbure_producer=self.owner).count() == 4
        assert lots.filter(carbure_producer=None, unknown_producer="").count() == 1
        assert lots.filter(carbure_producer=None, unknown_producer="Unknown").count() == 3

        assert lots.filter(carbure_production_site=self.owner_production_site).count() == 4
        assert lots.filter(carbure_production_site=None, unknown_production_site="Unknown Production Site").count() == 3

        assert lots.filter(carbure_supplier=self.owner).count() == 6
        assert lots.filter(carbure_supplier=None, unknown_supplier="UNKNOWN").count() == 2
        assert lots.filter(supplier_certificate=self.owner.default_certificate).count() == 6

        assert lots.filter(carbure_vendor=self.owner).count() == 0
        assert lots.filter(vendor_certificate=self.owner.default_certificate).count() == 0

        assert lots.filter(carbure_client=self.owner).count() == 1
        assert lots.filter(carbure_client=self.other).count() == 5

        assert lots.filter(carbure_delivery_site=self.owner_depot).count() == 1
        assert lots.filter(carbure_delivery_site=self.other_depot).count() == 5

    def test_invalid_dispatch_site_type_is_blocking(self):
        invalid_dispatch_site = Site.objects.create(
            name="Invalid Dispatch Site",
            site_type=Site.POWER_PLANT,
            country=self.FR,
        )
        lot = CarbureLot()

        errors = fill_dispatch_data(
            lot,
            {"carbure_dispatch_site_id": invalid_dispatch_site.id},
            {"sites": {invalid_dispatch_site.id: invalid_dispatch_site}, "countries": {}},
        )

        assert lot.carbure_dispatch_site is None
        assert len(errors) == 1
        assert errors[0].error == INVALID_DISPATCH_SITE
        assert errors[0].field == "dispatch_site"
        assert errors[0].is_blocking is True

    def test_biofuel_template_contains_dispatch_fields(self):
        self.owner.entity_type = Entity.PRODUCER
        self.owner.save()
        invalid_dispatch_site = Site.objects.create(
            name="Invalid Dispatch Site",
            site_type=Site.POWER_PLANT,
            country=self.FR,
        )

        template_path = template_v4(self.owner)
        workbook = openpyxl.load_workbook(template_path, read_only=True)
        headers = next(workbook["lots"].iter_rows(values_only=True))
        dispatch_sheet = workbook["SitesDExpedition"]
        dispatch_headers = next(dispatch_sheet.iter_rows(values_only=True))
        dispatch_rows = list(dispatch_sheet.iter_rows(min_row=2, values_only=True))
        workbook.close()

        assert headers[10:13] == ("dispatch_site", "dispatch_site_country", "dispatch_date")
        assert dispatch_headers == ("id", "name", "city", "country", "site_type")
        dispatch_ids = {row[0] for row in dispatch_rows}
        assert self.owner_production_site.id in dispatch_ids
        assert self.owner_depot.id in dispatch_ids
        assert invalid_dispatch_site.id not in dispatch_ids

    def test_lot_export_contains_dispatch_fields(self):
        lot = CarbureLot(
            dispatch_date=date(2025, 2, 15),
            carbure_dispatch_site=self.owner_production_site,
            dispatch_site_country=self.FR,
        )

        data = CarbureLotCSVSerializer(lot).data

        assert data["dispatch_date"] == "15/02/2025"
        assert data["dispatch_site"] == self.owner_production_site.name
        assert data["dispatch_site_country"] == "FR"

    def test_dispatch_fields_are_parsed_from_fixture(self):
        self.owner.entity_type = Entity.PRODUCER
        self.owner.has_trading = False
        self.owner.has_stocks = False
        self.owner.save()

        EntitySite.objects.update_or_create(entity=self.owner, site=self.owner_production_site)
        EntitySite.objects.filter(entity=self.owner, site=self.owner_depot).delete()

        lots = self.send_excel(self.owner, "test_lot_template_with_dispatch.xlsx")

        known_site_lot = lots.get(free_field="je suis le producteur")
        assert known_site_lot.carbure_dispatch_site_id == self.owner_production_site.id
        assert known_site_lot.unknown_dispatch_site is None
        assert known_site_lot.dispatch_site_country == self.FR
        assert known_site_lot.dispatch_date.isoformat() == "2025-02-15"

        unknown_site_lot = lots.get(free_field="je n'ai pas d'info de producteur")
        assert unknown_site_lot.carbure_dispatch_site is None
        assert unknown_site_lot.unknown_dispatch_site == "Unknown Dispatch Site"
        assert unknown_site_lot.dispatch_site_country.code_pays == "DE"
        assert unknown_site_lot.dispatch_date.isoformat() == "2025-02-16"

    def test_producer_trader_excel(self):
        self.owner.entity_type = Entity.PRODUCER
        self.owner.has_trading = True
        self.owner.has_stocks = True
        self.owner.save()

        # add production site, remove depot
        EntitySite.objects.update_or_create(entity=self.owner, site=self.owner_production_site)
        EntitySite.objects.filter(entity=self.owner, site=self.owner_depot).delete()

        lots = self.send_excel(self.owner, "test_lot_template.xlsx")

        assert lots.filter(carbure_producer=self.owner).count() == 4
        assert lots.filter(carbure_producer=None, unknown_producer="").count() == 1
        assert lots.filter(unknown_producer="Unknown").count() == 3

        assert lots.filter(carbure_production_site=self.owner_production_site).count() == 4
        assert lots.filter(unknown_production_site="Unknown Production Site").count() == 3

        assert lots.filter(carbure_supplier=self.owner).count() == 6
        assert lots.filter(carbure_supplier=None, unknown_supplier="Unknown").count() == 2
        assert lots.filter(supplier_certificate=self.owner.default_certificate).count() == 6

        assert lots.filter(carbure_vendor=self.owner).count() == 1
        assert lots.filter(vendor_certificate=self.owner.default_certificate).count() == 1

        assert lots.filter(carbure_client=self.owner).count() == 1
        assert lots.filter(carbure_client=self.other).count() == 5

        assert lots.filter(carbure_delivery_site=self.owner_depot).count() == 1
        assert lots.filter(carbure_delivery_site=self.other_depot).count() == 5

    def test_operator_excel(self):
        self.owner.entity_type = Entity.OPERATOR
        self.owner.has_trading = False
        self.owner.has_stocks = False
        self.owner.save()

        # remove production site, add depot
        EntitySite.objects.filter(entity=self.owner, site=self.owner_production_site).delete()
        EntitySite.objects.update_or_create(entity=self.owner, site=self.owner_depot)

        lots = self.send_excel(self.owner, "test_lot_template.xlsx")

        assert lots.filter(carbure_producer=None).count() == 8
        assert lots.filter(unknown_producer="Owner").count() == 3

        assert lots.filter(carbure_production_site=None).count() == 8
        assert lots.filter(unknown_production_site="Owner Production Site").count() == 4

        assert lots.filter(carbure_supplier=self.owner).count() == 6
        assert lots.filter(carbure_supplier=None, unknown_supplier="Unknown").count() == 2
        assert lots.filter(supplier_certificate=self.owner.default_certificate).count() == 6

        assert lots.filter(carbure_vendor=self.owner).count() == 1
        assert lots.filter(vendor_certificate=self.owner.default_certificate).count() == 1

        assert lots.filter(carbure_client=self.owner).count() == 2
        assert lots.filter(carbure_client=self.other).count() == 5

        assert lots.filter(carbure_delivery_site=self.owner_depot).count() == 1
        assert lots.filter(carbure_delivery_site=self.other_depot).count() == 5

    def test_trader_excel(self):
        self.owner.entity_type = Entity.TRADER
        self.owner.has_trading = True
        self.owner.has_stocks = True
        self.owner.save()

        # remove production site, add depot
        EntitySite.objects.filter(entity=self.owner, site=self.owner_production_site).delete()
        EntitySite.objects.update_or_create(entity=self.owner, site=self.owner_depot)

        lots = self.send_excel(self.owner, "test_lot_template.xlsx")

        assert lots.filter(carbure_producer=None).count() == 8
        assert lots.filter(unknown_producer="Owner").count() == 3

        assert lots.filter(carbure_production_site=None).count() == 8
        assert lots.filter(unknown_production_site="Owner Production Site").count() == 4

        assert lots.filter(carbure_supplier=self.owner).count() == 6
        assert lots.filter(carbure_supplier=None, unknown_supplier="Unknown").count() == 2
        assert lots.filter(supplier_certificate=self.owner.default_certificate).count() == 6

        assert lots.filter(carbure_vendor=self.owner).count() == 1
        assert lots.filter(vendor_certificate=self.owner.default_certificate).count() == 1

        assert lots.filter(carbure_client=self.owner).count() == 1
        assert lots.filter(carbure_client=self.other).count() == 5

        assert lots.filter(carbure_delivery_site=self.owner_depot).count() == 1
        assert lots.filter(carbure_delivery_site=self.other_depot).count() == 5

    def test_bulk_import_reads_usage_fields_from_excel(self):
        """Test that the bulk import correctly reads usage fields from the Excel file."""
        self.owner.entity_type = Entity.PRODUCER
        self.owner.has_trading = True
        self.owner.has_stocks = True
        self.owner.save()

        # add production site, remove depot
        EntitySite.objects.update_or_create(entity=self.owner, site=self.owner_production_site)
        EntitySite.objects.filter(entity=self.owner, site=self.owner_depot).delete()

        lots = self.send_excel(self.owner, "test_lot_template.xlsx")
        assert lots.filter(usage="").count() == 7

        lot_rfc = lots.filter(delivery_type="RFC")
        assert lot_rfc.count() == 1
        assert lot_rfc.first().usage == "OTHER"
        assert lot_rfc.first().usage_precision == "precisions"

    def test_excel_unknown_stock_id_returns_400(self):
        # When a row references a carbure_stock_id that does not exist,
        # the import should be rejected with a 400 and an explicit error message.
        filepath = f"{os.environ['CARBURE_HOME']}/web/transactions/fixtures/test_lot_template.xlsx"
        with open(filepath, "rb") as reader:
            file = SimpleUploadedFile("test_lot_template.xlsx", reader.read())

        with patch(
            "transactions.api.lots.add_excel.construct_carbure_lot",
            return_value=(None, []),
        ):
            response = self.client.post(
                reverse("transactions-lots-add-excel"),
                {"entity_id": self.owner.id, "file": file},
            )

        assert response.status_code == 400
        data = response.json()
        assert data["status"] == "error"
        # No lot should have been created (transaction is rolled back)
        assert CarbureLot.objects.filter(added_by=self.owner).count() == 0
