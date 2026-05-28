import io

from django.test import TestCase
from openpyxl import load_workbook

from biomethane.factories import (
    BiomethaneDigestateFactory,
    BiomethaneDigestateSpreadingFactory,
    BiomethaneProductionUnitFactory,
    BiomethaneSupplyInputFactory,
    BiomethaneSupplyPlanFactory,
)
from biomethane.factories.production_unit import BiomethaneDigestateStorageFactory
from biomethane.models import (
    BiomethaneContract,
    BiomethaneEnergy,
    BiomethaneEnergyMonthlyReport,
    BiomethaneProductionUnit,
)
from biomethane.services.declaration_export import _format_value, generate_annual_export
from core.models import Entity

EXPECTED_SHEETS = [
    "Contrat",
    "Unité de production",
    "Site d'injection",
    "Digestat",
    "Énergie",
    "Énergie mensuelle",
    "Approvisionnement",
    "Stockage digestat",
    "Épandage digestat",
]


class FormatValueTests(TestCase):
    """Unit tests for _format_value helper — no DB needed."""

    def test_true_formatted_as_oui(self):
        self.assertEqual(_format_value(True), "OUI")

    def test_false_formatted_as_non(self):
        self.assertEqual(_format_value(False), "NON")

    def test_none_formatted_as_empty_string(self):
        self.assertEqual(_format_value(None), "")


class GenerateAnnualExportTests(TestCase):
    """Integration tests for generate_annual_export service."""

    fixtures = ["json/countries.json"]
    YEAR = 2025

    def setUp(self):
        self.producer = Entity.objects.create(
            name="Test Producer",
            entity_type=Entity.BIOMETHANE_PRODUCER,
        )

    def _load_workbook(self, file):
        """Read the returned BufferedReader into an openpyxl workbook."""
        data = file.read()
        file.close()
        return load_workbook(io.BytesIO(data))

    def _find_value_by_label(self, worksheet, label):
        for row_idx in range(2, worksheet.max_row + 1):
            if worksheet.cell(row=row_idx, column=1).value == label:
                return worksheet.cell(row=row_idx, column=2).value
        return None

    def test_returns_readable_file(self):
        """generate_annual_export returns a readable BufferedReader."""
        from io import BufferedReader

        file = generate_annual_export(self.producer, self.YEAR)
        self.assertIsInstance(file, BufferedReader)
        self.assertTrue(file.readable())
        file.close()

    def test_excel_has_9_sheets(self):
        """Generated workbook contains exactly the 9 expected sheets."""
        file = generate_annual_export(self.producer, self.YEAR)
        wb = self._load_workbook(file)
        self.assertEqual(wb.sheetnames, EXPECTED_SHEETS)

    # Test 12
    def test_all_sheets_have_headers(self):
        """Every sheet has a non-empty first row even without data."""
        file = generate_annual_export(self.producer, self.YEAR)
        wb = self._load_workbook(file)
        for sheet_name in EXPECTED_SHEETS:
            with self.subTest(sheet=sheet_name):
                first_row_values = [cell.value for cell in wb[sheet_name][1]]
                self.assertTrue(
                    any(v is not None for v in first_row_values),
                    f"Sheet '{sheet_name}' has an empty first row",
                )

    def test_contract_sheet_contains_data(self):
        """'Contrat' sheet has data rows when a contract exists for the producer."""
        BiomethaneContract.objects.create(producer=self.producer)
        file = generate_annual_export(self.producer, self.YEAR)
        wb = self._load_workbook(file)
        self.assertGreater(wb["Contrat"].max_row, 1)

    def test_supply_inputs_sheet_rows(self):
        """'Approvisionnement' sheet has one row per BiomethaneSupplyInput."""
        supply_plan = BiomethaneSupplyPlanFactory.create(producer=self.producer, year=self.YEAR)
        BiomethaneSupplyInputFactory.create(supply_plan=supply_plan)
        BiomethaneSupplyInputFactory.create(supply_plan=supply_plan)
        BiomethaneSupplyInputFactory.create(supply_plan=supply_plan)

        file = generate_annual_export(self.producer, self.YEAR)
        wb = self._load_workbook(file)
        # 1 header row + 3 data rows
        self.assertEqual(wb["Approvisionnement"].max_row, 4)

    def test_digestate_spreading_rows(self):
        """'Épandage digestat' sheet has one row per BiomethaneDigestateSpreading."""
        digestate = BiomethaneDigestateFactory.create(producer=self.producer, year=self.YEAR)
        BiomethaneDigestateSpreadingFactory.create(digestate=digestate)
        BiomethaneDigestateSpreadingFactory.create(digestate=digestate)

        file = generate_annual_export(self.producer, self.YEAR)
        wb = self._load_workbook(file)
        # 1 header row + 2 data rows
        self.assertEqual(wb["Épandage digestat"].max_row, 3)

    def test_storage_sheet_rows(self):
        """'Stockage digestat' sheet has one row per BiomethaneDigestateStorage."""
        BiomethaneDigestateStorageFactory.create(producer=self.producer)
        BiomethaneDigestateStorageFactory.create(producer=self.producer)

        file = generate_annual_export(self.producer, self.YEAR)
        wb = self._load_workbook(file)
        # 1 header row + 2 data rows
        self.assertEqual(wb["Stockage digestat"].max_row, 3)

    def test_energy_monthly_reports_sheet_rows(self):
        """'Énergie mensuelle' sheet has one row per BiomethaneEnergyMonthlyReport."""
        energy = BiomethaneEnergy.objects.create(producer=self.producer, year=self.YEAR)
        BiomethaneEnergyMonthlyReport.objects.create(
            energy=energy,
            month=1,
            injected_volume_nm3=100.0,
            average_monthly_flow_nm3_per_hour=10.0,
        )
        BiomethaneEnergyMonthlyReport.objects.create(
            energy=energy,
            month=2,
            injected_volume_nm3=200.0,
            average_monthly_flow_nm3_per_hour=20.0,
        )

        file = generate_annual_export(self.producer, self.YEAR)
        wb = self._load_workbook(file)
        # 1 header row + 2 data rows
        self.assertEqual(wb["Énergie mensuelle"].max_row, 3)
        self.assertEqual(wb["Énergie mensuelle"]["A2"].value, "Janvier")
        self.assertEqual(wb["Énergie mensuelle"]["A3"].value, "Février")
        self.assertEqual(wb["Énergie mensuelle"]["D1"].value, "Heures d'injection (h)")
        self.assertEqual(wb["Énergie mensuelle"]["D2"].value, 10.0)
        self.assertEqual(wb["Énergie mensuelle"]["D3"].value, 10.0)

    def test_contract_json_list_choices_are_exported_with_labels(self):
        BiomethaneContract.objects.create(
            producer=self.producer,
            has_complementary_investment_aid=True,
            complementary_aid_organisms=[
                BiomethaneContract.COMPLEMENTARY_AID_ORGANISM_ADEME,
                BiomethaneContract.COMPLEMENTARY_AID_ORGANISM_OTHER,
            ],
        )

        file = generate_annual_export(self.producer, self.YEAR)
        wb = self._load_workbook(file)

        value = self._find_value_by_label(wb["Contrat"], "Aide complémentaire attribuée par")
        self.assertEqual(value, "Ademe, Autre")

    def test_production_unit_choice_is_exported_with_label(self):
        BiomethaneProductionUnitFactory.create(
            producer=self.producer,
            methanization_process=BiomethaneProductionUnit.PLUG_FLOW_SEMI_CONTINUOUS,
        )

        file = generate_annual_export(self.producer, self.YEAR)
        wb = self._load_workbook(file)

        value = self._find_value_by_label(wb["Unité de production"], "Procédé méthanisation")
        self.assertEqual(value, "En piston (semi-continu)")

    def test_empty_producer_has_only_headers(self):
        """When no data exists, all sheets contain only the header row."""
        file = generate_annual_export(self.producer, self.YEAR)
        wb = self._load_workbook(file)
        for sheet_name in EXPECTED_SHEETS:
            with self.subTest(sheet=sheet_name):
                self.assertEqual(
                    wb[sheet_name].max_row,
                    1,
                    f"Sheet '{sheet_name}' should have only the header row",
                )

    def test_wrong_year_year_filtered_sheets_are_empty(self):
        """Year-filtered sheets are empty when data exists for a different year."""
        BiomethaneDigestateFactory.create(producer=self.producer, year=self.YEAR)
        energy = BiomethaneEnergy.objects.create(producer=self.producer, year=self.YEAR)
        BiomethaneEnergyMonthlyReport.objects.create(
            energy=energy,
            month=1,
            injected_volume_nm3=100.0,
            average_monthly_flow_nm3_per_hour=10.0,
        )
        supply_plan = BiomethaneSupplyPlanFactory.create(producer=self.producer, year=self.YEAR)
        BiomethaneSupplyInputFactory.create(supply_plan=supply_plan)

        file = generate_annual_export(self.producer, self.YEAR - 1)
        wb = self._load_workbook(file)

        for sheet_name in ("Digestat", "Énergie", "Énergie mensuelle", "Approvisionnement"):
            with self.subTest(sheet=sheet_name):
                self.assertEqual(
                    wb[sheet_name].max_row,
                    1,
                    f"Sheet '{sheet_name}' should be empty when exporting a different year",
                )
