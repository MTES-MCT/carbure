from django.test import TestCase
from openpyxl import load_workbook

from core.models import Entity
from core.models.certificate import GenericCertificate
from h2.handlers import H2ActionHandler
from h2.tests.excel import filled_h2_template
from traceability.factories import MaterialFactory
from traceability.services.action_excel import build_action_import_template, parse_action_import_file
from transactions.factories.certificate import GenericCertificateFactory
from transactions.models import Site


class H2ActionExcelTemplateTest(TestCase):
    def test_template_uses_handler_columns_and_lookup_options(self):
        entity = Entity.objects.create(name="HRS", entity_type=Entity.HRS)
        hydrogen = MaterialFactory(code="H2-GASE", name="Hydrogène gazeux")
        MaterialFactory(code="BIO-WOOD", name="Bois")
        station = Site.objects.create(
            name="Station Paris",
            site_type=Site.H2_REFUELING_STATION,
            created_by=entity,
        )
        Site.objects.create(name="Dépôt Lyon", site_type=Site.EFS, created_by=entity)
        Site.objects.create(
            name="Station Lyon",
            site_type=Site.H2_REFUELING_STATION,
            created_by=Entity.objects.create(name="Other HRS", entity_type=Entity.HRS),
        )
        certifhy = GenericCertificateFactory.create(
            certificate_id="CHY-001",
            certificate_type=GenericCertificate.CERTIFHY,
        )
        GenericCertificateFactory.create(
            certificate_id="ISCC-001",
            certificate_type=GenericCertificate.ISCC,
        )

        file_handle = build_action_import_template(H2ActionHandler(), entity)
        workbook = load_workbook(filename=file_handle)
        file_handle.close()
        self.addCleanup(workbook.close)

        headers = [cell.value for cell in workbook["Import actions H2"][1]]
        references = workbook["References"]

        self.assertEqual(headers, [column["header"] for column in H2ActionHandler.excel_columns])

        def reference_values(key):
            index = next(i for i, spec in enumerate(H2ActionHandler.excel_columns, start=1) if spec.get("key") == key)
            return [
                references.cell(row=row, column=index).value
                for row in range(2, references.max_row + 1)
                if references.cell(row=row, column=index).value is not None
            ]

        self.assertIn(hydrogen.name, reference_values("material"))
        self.assertNotIn("Bois", reference_values("material"))
        self.assertIn(station.name, reference_values("site"))
        self.assertNotIn("Dépôt Lyon", reference_values("site"))
        self.assertNotIn("Station Lyon", reference_values("site"))
        self.assertIn(certifhy.certificate_id, reference_values("certificate"))
        self.assertNotIn("ISCC-001", reference_values("certificate"))


class ParseH2ActionImportFileTest(TestCase):
    def test_maps_h2_extra_columns(self):
        buffer = filled_h2_template(
            pos_id="H2-001",
            material_name="Hydrogène gazeux",
            site_name="Station Paris",
        )

        rows = parse_action_import_file(buffer, H2ActionHandler())

        self.assertEqual(rows[0]["lot_id"], "LOT-001")
        self.assertIn("etd1", rows[0])
        self.assertIn("etd2", rows[0])
        self.assertNotIn("etd", rows[0])
