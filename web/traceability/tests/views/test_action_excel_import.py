from datetime import date
from decimal import Decimal
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from core.models import Entity
from core.tests_utils import setup_current_user
from h2.factories.h2_station import H2StationFactory
from h2.handlers import H2ActionHandler
from traceability.factories import MaterialFactory
from traceability.models import Action
from traceability.models.action_status import ActionStatus
from traceability.services.action_excel import build_action_import_template, parse_action_import_file
from transactions.models.site import Site


def filled_h2_template(*, pos_id, material_name, site_name, shipping_date=date(2026, 1, 15), extra_headers=None):
    file_handle = build_action_import_template(H2ActionHandler())
    workbook = load_workbook(filename=BytesIO(file_handle.read()))
    file_handle.close()
    sheet = workbook["Import actions H2"]
    values = [
        pos_id,
        material_name,
        Decimal("120000.000"),
        site_name,
        shipping_date,
        25,
        Action.ROAD,
        0,
        0,
        0,
        0,
        0,
    ]
    for column, value in enumerate(values, start=1):
        sheet.cell(row=2, column=column, value=value)
    if extra_headers:
        start = len(H2ActionHandler.excel_columns) + 1
        for offset, (header, value) in enumerate(extra_headers):
            sheet.cell(row=1, column=start + offset, value=header)
            sheet.cell(row=2, column=start + offset, value=value)

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    buffer.seek(0)
    return buffer


class ParseActionImportFileTest(TestCase):
    def test_maps_headers_to_keys_and_ignores_unknown_columns(self):
        buffer = filled_h2_template(
            pos_id="H2-001",
            material_name="Hydrogène gazeux",
            site_name="Station Paris",
            extra_headers=[("Producteur", "Air Liquide")],
        )

        rows = parse_action_import_file(buffer, H2ActionHandler())

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["pos_id"], "H2-001")
        self.assertEqual(rows[0]["material"], "Hydrogène gazeux")
        self.assertNotIn("Producteur", rows[0])
        self.assertNotIn("producer", rows[0])


class ActionExcelImportViewTest(APITestCase):
    fixtures = ["json/countries.json"]

    def setUp(self):
        self.entity = Entity.objects.create(name="HRS", entity_type=Entity.HRS)
        setup_current_user(self, "tester@carbure.local", "Tester", "password", [(self.entity, "RW")])
        self.url = reverse("traceability-action-import-actions")
        self.material = MaterialFactory(code="H2-GASE", name="Hydrogène gazeux")
        self.station = H2StationFactory.create(
            created_by=self.entity,
        )
        self.non_h2_station = Site.objects.create(
            name="EFS",
            site_type=Site.EFS,
            created_by=self.entity,
        )

    def _post(self, buffer):
        uploaded = SimpleUploadedFile(
            "import.xlsx",
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        return self.client.post(
            self.url,
            {"file": uploaded},
            query_params={"entity_id": self.entity.id, "industry": Action.H2},
            format="multipart",
        )

    def test_import_creates_an_init_action_for_the_entity(self):
        response = self._post(
            filled_h2_template(
                pos_id="H2-001",
                material_name=self.material.name,
                site_name=self.station.name,
            )
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["rows_imported"], 1)

        action = Action.objects.get(pos_id="H2-001")
        self.assertEqual(action.holder, self.entity)
        self.assertEqual(action.industry, Action.H2)
        self.assertEqual(action.type, Action.INIT)
        self.assertEqual(action.material, self.material)
        self.assertEqual(action.site.id, self.station.id)
        self.assertEqual(action.shipping_method, Action.ROAD)
        self.assertEqual(action.working_date, date(2026, 1, 15))
        self.assertEqual(action.action_statuses.get().status, ActionStatus.CREATED)

    def test_import_accepts_shipping_date_as_day_month_year(self):
        response = self._post(
            filled_h2_template(
                pos_id="H2-001",
                material_name=self.material.name,
                site_name=self.station.name,
                shipping_date="15/01/2026",
            )
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(Action.objects.get(pos_id="H2-001").shipping_date, date(2026, 1, 15))

    def test_import_rejects_unknown_material(self):
        response = self._post(
            filled_h2_template(
                pos_id="H2-001",
                material_name="Bois",
                site_name=self.station.name,
            )
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(Action.objects.count(), 0)
        self.assertTrue(response.data["validation_errors"])

    def test_import_rejects_material_from_another_industry(self):
        MaterialFactory(code="BIO-WOOD", name="Bois")

        response = self._post(
            filled_h2_template(
                pos_id="H2-001",
                material_name="Bois",
                site_name=self.station.name,
            )
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(Action.objects.count(), 0)
        self.assertTrue(response.data["validation_errors"])

    def test_import_rejects_site_from_another_industry(self):
        response = self._post(
            filled_h2_template(
                pos_id="H2-001",
                material_name=self.material.name,
                site_name=self.non_h2_station.name,
            )
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(Action.objects.count(), 0)
        self.assertTrue(response.data["validation_errors"])
