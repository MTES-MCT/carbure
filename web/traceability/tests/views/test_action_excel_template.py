from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Entity
from core.tests_utils import setup_current_user
from traceability.models import Action


class ActionExcelTemplateViewTest(TestCase):
    def setUp(self):
        self.entity = Entity.objects.create(name="HRS", entity_type=Entity.HRS)
        setup_current_user(self, "tester@carbure.local", "Tester", "password", [(self.entity, "RW")])
        self.url = reverse("traceability-action-download-import-template")

    def test_download_returns_the_h2_excel_template(self):
        response = self.client.get(self.url, {"entity_id": self.entity.id, "industry": Action.H2})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.ms-excel")

        workbook = load_workbook(filename=BytesIO(response.content))
        self.addCleanup(workbook.close)
        self.assertIn("Import actions H2", workbook.sheetnames)
