from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from core.import_export_template import get_data_start_row
from traceability.handlers.action import ActionIndustryHandler
from traceability.models import Action
from traceability.services.action_excel import build_action_import_template


class GenericExcelHandler(ActionIndustryHandler):
    industry = Action.H2


GENERIC_ROW = {
    "pos_id": "ACT-001",
    "material": "Matière",
    "certificate": "CERT-001",
    "quantity": Decimal("120000.000"),
    "site": "Site",
    "shipping_date": date(2026, 1, 15),
    "shipping_distance": 25,
    "shipping_method": "Transport routier",
    "working_date": date(2026, 2, 1),
    "eec": 0,
    "el": 0,
    "ei": 0,
    "ep": 0,
    "etd": 0,
    "eu": 0,
    "eccs": 0,
    "esca": 0,
    "eccr": 0,
}


def fill_action_template(handler, values_by_key, extra_headers=None):
    file_handle = build_action_import_template(handler)
    workbook = load_workbook(filename=BytesIO(file_handle.read()))
    file_handle.close()
    sheet = workbook[f"Import actions {handler.industry}"]
    data_row = get_data_start_row(handler.excel_columns)
    for column, spec in enumerate(handler.excel_columns, start=1):
        key = spec.get("key")
        if key:
            sheet.cell(row=data_row, column=column, value=values_by_key[key])
    if extra_headers:
        start = len(handler.excel_columns) + 1
        for offset, (header, value) in enumerate(extra_headers):
            sheet.cell(row=1, column=start + offset, value=header)
            sheet.cell(row=data_row, column=start + offset, value=value)

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    buffer.seek(0)
    return buffer


def filled_generic_template(*, material_name, site_name, certificate_id="", extra_headers=None, **overrides):
    return fill_action_template(
        GenericExcelHandler(),
        {
            **GENERIC_ROW,
            "material": material_name,
            "site": site_name,
            "certificate": certificate_id,
            **overrides,
        },
        extra_headers=extra_headers,
    )
