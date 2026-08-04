import os
from unittest.mock import Mock, patch

import openpyxl
from django.test import TestCase

from core.models import Entity
from entity.factories import EntityFactory
from tiruert.services.operation_excel_template import (
    ENTITIES_SHEET_NAME,
    FIRST_DATA_ROW,
    KEY_ROW,
    MAIN_SHEET_NAME,
    TABLE_HEADERS,
    _get_operator_lots,
    create_operation_import_template,
    get_tiruert_operator_queryset,
)


class OperationExcelTemplateServiceTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.allowed_entity_1 = EntityFactory.create(
            name="Allowed 1",
            entity_type=Entity.OPERATOR,
            is_enabled=True,
            is_tiruert_liable=True,
            accise_number="ACC-001",
        )
        cls.allowed_entity_2 = EntityFactory.create(
            name="Allowed 2",
            entity_type=Entity.OPERATOR,
            is_enabled=True,
            is_tiruert_liable=True,
            accise_number="ACC-002",
        )
        cls.disabled_entity = EntityFactory.create(
            name="Disabled",
            entity_type=Entity.OPERATOR,
            is_enabled=False,
            is_tiruert_liable=True,
            accise_number="ACC-003",
        )
        cls.no_accise_entity = EntityFactory.create(
            name="No Accise",
            entity_type=Entity.OPERATOR,
            is_enabled=True,
            is_tiruert_liable=True,
            accise_number="",
        )
        cls.not_liable_entity = EntityFactory.create(
            name="Not Liable",
            entity_type=Entity.OPERATOR,
            is_enabled=True,
            is_tiruert_liable=False,
            accise_number="ACC-004",
        )

    def test_get_tiruert_operator_queryset_filters_and_orders_entities(self):
        """Should return only enabled, liable operator entities with an accise number, ordered by name."""
        names = list(get_tiruert_operator_queryset().values_list("name", flat=True))

        self.assertEqual(names, ["Allowed 1", "Allowed 2"])

    @patch("tiruert.services.operation_excel_template.CarbureLot.objects.filter")
    @patch("tiruert.services.operation_excel_template.BalanceService.calculate_balance")
    @patch("tiruert.services.operation_excel_template.Operation.objects.filter")
    def test_get_operator_lots_returns_positive_balances_sorted_by_carbure_id(
        self, mock_operation_filter, mock_calculate_balance, mock_lot_filter
    ):
        """Should keep only positive balances, enrich metadata, and sort lots by carbure_id."""
        entity_id = 42
        mock_operation_filter.return_value = Mock(name="operations_queryset")
        mock_calculate_balance.return_value = {
            ("ESSENCE", "CONV", "ETH", 1): {"available_balance": 12.345, "emission_rate_per_mj": 7.89},
            ("ESSENCE", "CONV", "EMAG", 2): {"available_balance": 0, "emission_rate_per_mj": 3.21},
            ("GAZOLE", "ANN-IX-A", "HVO", 3): {"available_balance": 3.0, "emission_rate_per_mj": 1.23},
        }
        mock_lot_filter.return_value.values_list.return_value = [
            (1, "CARB-2", "Maize"),
            (3, "CARB-1", "Wheat"),
        ]

        lots = _get_operator_lots(entity_id)

        self.assertEqual([lot["carbure_id"] for lot in lots], ["CARB-1", "CARB-2"])
        self.assertEqual(lots[0]["lot_id"], 3)
        self.assertEqual(lots[0]["available_volume"], 3.0)
        self.assertEqual(lots[0]["feedstock"], "Wheat")
        self.assertEqual(lots[1]["lot_id"], 1)
        self.assertEqual(lots[1]["available_volume"], 12.35)
        mock_operation_filter.assert_called_once()
        mock_calculate_balance.assert_called_once()
        mock_lot_filter.assert_called_once_with(id__in=[1, 2, 3])

    @patch("tiruert.services.operation_excel_template._get_operator_lots")
    @patch("tiruert.services.operation_excel_template.get_tiruert_operator_queryset")
    def test_create_operation_import_template_builds_expected_workbook(self, mock_get_entities, mock_get_lots):
        """Should generate a workbook with the main and entities sheets wired as expected."""
        mock_get_entities.return_value = [self.allowed_entity_1, self.allowed_entity_2]
        mock_get_lots.return_value = [
            {
                "lot_id": 3,
                "carbure_id": "CARB-1",
                "available_volume": 3.0,
                "emission_rate_per_mj": 1.23,
                "biofuel": "HVO",
                "customs_category": "ANN-IX-A",
                "feedstock": "Wheat",
            },
            {
                "lot_id": 1,
                "carbure_id": "CARB-2",
                "available_volume": 12.35,
                "emission_rate_per_mj": 7.89,
                "biofuel": "ETH",
                "customs_category": "CONV",
                "feedstock": "Maize",
            },
        ]

        fh = create_operation_import_template(self.allowed_entity_1.id)
        self.addCleanup(lambda: fh.close())
        self.addCleanup(lambda: os.path.exists(fh.name) and os.remove(fh.name))

        workbook = openpyxl.load_workbook(fh.name)
        main_sheet = workbook[MAIN_SHEET_NAME]
        entities_sheet = workbook[ENTITIES_SHEET_NAME]

        self.assertEqual(workbook.sheetnames, [MAIN_SHEET_NAME, ENTITIES_SHEET_NAME])
        self.assertEqual(main_sheet[1][0].value, TABLE_HEADERS[0][0])
        self.assertEqual(main_sheet[1][len(TABLE_HEADERS) - 1].value, TABLE_HEADERS[-1][0])
        self.assertEqual(main_sheet[2][0].value, TABLE_HEADERS[0][1])
        self.assertEqual(main_sheet[3][0].value, "CARB-1")
        self.assertEqual(main_sheet[3][1].value, 3)
        self.assertEqual(main_sheet[3][7].value, None)
        self.assertEqual(main_sheet[4][0].value, "CARB-2")
        self.assertEqual(main_sheet[4][1].value, 1)
        self.assertEqual(entities_sheet[2][0].value, self.allowed_entity_1.name)
        self.assertEqual(entities_sheet[3][0].value, self.allowed_entity_2.name)
        self.assertEqual(entities_sheet[2][1].value, self.allowed_entity_1.id)
        self.assertEqual(entities_sheet[3][1].value, self.allowed_entity_2.id)
        self.assertEqual(main_sheet.sheet_state, "visible")
        self.assertEqual(entities_sheet.sheet_state, "hidden")
        self.assertGreaterEqual(len(main_sheet.data_validations.dataValidation), 4)
        self.assertEqual(main_sheet.row_dimensions[KEY_ROW + 1].hidden, True)
        self.assertEqual(main_sheet.column_dimensions["B"].hidden, True)
        self.assertEqual(main_sheet.column_dimensions["K"].hidden, True)
        self.assertEqual(main_sheet.column_dimensions["B"].hidden, True)
        self.assertEqual(main_sheet.cell(row=FIRST_DATA_ROW + 1, column=10).value, None)
