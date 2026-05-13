from datetime import date
from io import BytesIO

import openpyxl
from django.test import TestCase

from core.models import Entity, UserRights
from core.tests_utils import PermissionTestMixin, setup_current_user
from tiruert.models import FossilFuel, FossilFuelCategory, MacFossilFuel
from tiruert.permissions import HasTiruertRightsObjectives, HasTiruertWriteRights
from tiruert.views import MacFossilFuelViewSet


class MacFossilFuelExportEndpointSecurityTest(TestCase):
    def setUp(self):
        super().setUp()
        self.url = "/api/tiruert/mac-fossil-fuel/export/"

        self.allowed_entity = Entity.objects.create(
            name="Operateur autorise test",
            entity_type=Entity.OPERATOR,
            is_tiruert_liable=True,
            accise_number="ACC-ALLOWED",
            has_mac=False,
        )
        self.other_entity = Entity.objects.create(
            name="Operateur autre societe test",
            entity_type=Entity.OPERATOR,
            is_tiruert_liable=True,
            accise_number="ACC-OTHER",
            has_mac=False,
        )

        category = FossilFuelCategory.objects.create(name="Essence", pci_litre=32.0)
        self.fuel = FossilFuel.objects.create(
            label="SP95",
            nomenclature="SP95",
            fuel_category=category,
            pci_litre=32.0,
            masse_volumique=0.75,
        )

        MacFossilFuel.objects.create(
            fuel=self.fuel,
            operator=self.allowed_entity,
            volume=100.0,
            period=1,
            year=2023,
            start_date=date(2023, 1, 1),
            end_date=date(2023, 1, 31),
        )
        MacFossilFuel.objects.create(
            fuel=self.fuel,
            operator=self.other_entity,
            volume=200.0,
            period=1,
            year=2023,
            start_date=date(2023, 1, 1),
            end_date=date(2023, 1, 31),
        )

        setup_current_user(
            self,
            "tiruert-user@carbure.local",
            "Tiruert User",
            "secret",
            [(self.allowed_entity, UserRights.ADMIN)],
        )

    def test_anonymous_user_cannot_export_mac_fossil_fuel(self):
        self.client.logout()

        response = self.client.get(self.url, {"entity_id": self.allowed_entity.id})

        self.assertEqual(response.status_code, 403)

    def test_authenticated_user_cannot_export_mac_from_other_entity(self):
        response = self.client.get(self.url, {"entity_id": self.other_entity.id})

        self.assertEqual(response.status_code, 403)

    def test_authenticated_user_can_export_own_entity_macs(self):
        MacFossilFuel.objects.create(
            fuel=self.fuel,
            operator=self.allowed_entity,
            volume=50.0,
            period=2,
            year=2023,
            start_date=date(2023, 2, 1),
            end_date=date(2023, 2, 28),
        )

        response = self.client.get(self.url, {"entity_id": self.allowed_entity.id, "year": 2023})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        workbook = openpyxl.load_workbook(BytesIO(response.content))
        sheet = workbook.active

        # Header + 1 ligne de MAC agrégée pour l'entité autorisée
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(row=2, column=2).value, self.allowed_entity.name)
        self.assertEqual(sheet.cell(row=2, column=4).value, 150.0)
        self.assertEqual(sheet.cell(row=2, column=5).value, 2023)

    def test_authenticated_user_can_list_own_entity_macs(self):
        response = self.client.get("/api/tiruert/mac-fossil-fuel/", {"entity_id": self.allowed_entity.id, "year": 2023})

        self.assertEqual(response.status_code, 200)
        results = response.json()["results"]
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["operator"], self.allowed_entity.name)
        self.assertEqual(results[0]["volume"], 100.0)

    def test_authenticated_user_can_replace_own_entity_macs(self):
        response = self.client.put(
            "/api/tiruert/mac-fossil-fuel/replace/",
            data=[
                {"fuel": "SP95", "month": 1, "volume": 300.0},
                {"fuel": "SP95", "month": 2, "volume": 400.0},
            ],
            content_type="application/json",
            QUERY_STRING=f"entity_id={self.allowed_entity.id}&year=2023",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(MacFossilFuel.objects.filter(operator=self.allowed_entity, year=2023).count(), 2)
        self.assertTrue(MacFossilFuel.objects.filter(operator=self.other_entity, year=2023, volume=200.0).exists())
        self.assertEqual(
            list(
                MacFossilFuel.objects.filter(operator=self.allowed_entity, year=2023)
                .order_by("period")
                .values_list("period", "volume")
            ),
            [(202301, 300.0), (202302, 400.0)],
        )


class MacFossilFuelExportViewSetPermissionsTest(TestCase, PermissionTestMixin):
    def test_mac_fossil_fuel_export_uses_objectives_permission(self):
        self.assertViewPermissions(
            MacFossilFuelViewSet,
            [
                (
                    ["export_macfossilfuel_to_excel", "list"],
                    [HasTiruertRightsObjectives()],
                ),
                (
                    ["replace"],
                    [HasTiruertWriteRights()],
                ),
            ],
        )
