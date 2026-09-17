from datetime import date
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.request import Request
from rest_framework.test import APIRequestFactory

from core.models import Biocarburant, DeclarationPeriod, Entity, MatierePremiere
from core.tests_utils import setup_current_user
from core.utils import truncate
from tiruert.models import Operation, OperationDetail
from tiruert.views.operation.operation import OperationViewSet
from transactions.factories import CarbureLotFactory
from transactions.models import Depot


class OperationViewSetIntegrationTest(TestCase):
    """Test OperationViewSet CRUD operations."""

    fixtures = [
        "json/biofuels.json",
        "json/feedstock.json",
        "json/countries.json",
        "json/entities.json",
        "json/depots.json",
        "json/entities_sites.json",
    ]

    @classmethod
    def setUpTestData(cls):
        """Create test data once for the entire test class."""
        # Create declaration period for current year (covers entire year for tests)
        current_year = date.today().year
        DeclarationPeriod.objects.create(
            year=current_year,
            start_date=date(current_year, 1, 1),
            end_date=date(current_year, 12, 31),
            app=DeclarationPeriod.TIRUERT,
        )

        # Get references from fixtures
        cls.entity = Entity.objects.filter(entity_type=Entity.OPERATOR).first()
        cls.other_entity = Entity.objects.filter(entity_type=Entity.OPERATOR).exclude(id=cls.entity.id).first()
        cls.depot = Depot.objects.first()
        cls.feedstock_conv = MatierePremiere.biofuel.filter(category="CONV").first()
        cls.biofuel_eth = Biocarburant.objects.get(code="ETH")

        # Create TIRUERT operations directly for testing ViewSet
        # Operation 1: INCORPORATION (VALIDATED)
        cls.operation_incorporation = Operation.objects.create(
            type=Operation.INCORPORATION,
            status=Operation.VALIDATED,
            customs_category=MatierePremiere.CONV,
            biofuel=cls.biofuel_eth,
            credited_entity=cls.entity,
            to_depot=cls.depot,
            renewable_energy_share=1.0,
        )

        # Operation 2: MAC_BIO (VALIDATED)
        cls.operation_mac_bio = Operation.objects.create(
            type=Operation.MAC_BIO,
            status=Operation.VALIDATED,
            customs_category=MatierePremiere.CONV,
            biofuel=cls.biofuel_eth,
            credited_entity=cls.entity,
            to_depot=cls.depot,
            renewable_energy_share=1.0,
        )

        # Operation 3: LIVRAISON_DIRECTE (VALIDATED)
        cls.operation_livraison = Operation.objects.create(
            type=Operation.LIVRAISON_DIRECTE,
            status=Operation.VALIDATED,
            customs_category=MatierePremiere.CONV,
            biofuel=cls.biofuel_eth,
            credited_entity=cls.entity,
            to_depot=cls.depot,
            renewable_energy_share=1.0,
        )

        # Operation 4: CESSION (PENDING)
        cls.operation_cession = Operation.objects.create(
            type=Operation.CESSION,
            status=Operation.PENDING,
            customs_category=MatierePremiere.CONV,
            biofuel=cls.biofuel_eth,
            credited_entity=cls.other_entity,
            debited_entity=cls.entity,
            to_depot=cls.depot,
            renewable_energy_share=1.0,
        )

        # Add OperationDetails to first operation for testing details endpoint
        cls.lot1 = CarbureLotFactory.create(
            carbure_client=cls.entity,
            feedstock=cls.feedstock_conv,
            biofuel=cls.biofuel_eth,
            lot_status="ACCEPTED",
            delivery_type="BLENDING",
            volume=1000,
            ghg_total=10.5,
        )
        cls.lot2 = CarbureLotFactory.create(
            carbure_client=cls.entity,
            feedstock=cls.feedstock_conv,
            biofuel=cls.biofuel_eth,
            lot_status="ACCEPTED",
            delivery_type="BLENDING",
            volume=2000,
            ghg_total=12.3,
        )

        # Create OperationDetails
        OperationDetail.objects.create(
            operation=cls.operation_incorporation, lot=cls.lot1, volume=500, emission_rate_per_mj=10.5
        )
        OperationDetail.objects.create(
            operation=cls.operation_incorporation, lot=cls.lot2, volume=1500, emission_rate_per_mj=12.3
        )

    def setUp(self):
        super().setUp()
        self.user = setup_current_user(self, "tester@carbure.local", "Tester", "gogogo", [(self.entity, "ADMIN")])
        self.url = reverse("operations-list")
        self.factory = APIRequestFactory()

    def test_list_operations_returns_all_operations(self):
        """Test GET /operations/ returns list of operations."""
        query = {"entity_id": self.entity.id}
        response = self.client.get(self.url, query)

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("count", data)
        self.assertIn("results", data)
        self.assertEqual(data["count"], 4)

    def test_list_operations_returns_gpl_sector_for_gpl_compatible_biofuel(self):
        """The list annotation should expose GPL_C for GPL-compatible biofuels."""
        biofuel_gpl = Biocarburant.objects.create(
            code="TEST-GPL-LIST",
            name="Test GPL",
            name_en="Test GPL",
            description="Test GPL",
            pci_litre=24,
            compatible_essence=False,
            compatible_diesel=False,
            compatible_gpl=True,
        )
        operation = Operation.objects.create(
            type=Operation.INCORPORATION,
            status=Operation.VALIDATED,
            customs_category=MatierePremiere.CONV,
            biofuel=biofuel_gpl,
            credited_entity=self.entity,
            to_depot=self.depot,
            renewable_energy_share=1.0,
        )

        response = self.client.get(self.url, {"entity_id": self.entity.id})

        self.assertEqual(response.status_code, 200)
        result = next(item for item in response.json()["results"] if item["id"] == operation.id)
        self.assertEqual(result["sector"], Operation.GPL_C)

    def test_list_operations_filters_teneur_by_objective_sector(self):
        """The sector filter should use objective_sector for TENEUR operations."""
        operation = Operation.objects.create(
            type=Operation.TENEUR,
            status=Operation.VALIDATED,
            customs_category=MatierePremiere.CONV,
            biofuel=self.biofuel_eth,
            debited_entity=self.entity,
            objective_sector=Operation.ESSENCE,
            renewable_energy_share=1.0,
        )

        response = self.client.get(
            self.url,
            {"entity_id": self.entity.id, "sector": Operation.ESSENCE},
        )

        self.assertEqual(response.status_code, 200)
        result = next(item for item in response.json()["results"] if item["id"] == operation.id)
        self.assertEqual(result["sector"], Operation.ESSENCE)

    def test_list_queryset_clears_details_prefetch_when_details_not_requested(self):
        """List queryset should not keep the default details prefetch unless details are requested."""
        django_request = self.factory.get(self.url, {"entity_id": self.entity.id})
        django_request.entity = self.entity
        django_request.unit = "l"

        view = OperationViewSet()
        view.request = Request(django_request)
        view.action = "list"

        queryset = view.get_queryset()

        self.assertEqual(queryset._prefetch_related_lookups, ())

    def test_list_queryset_prefetches_details_when_requested(self):
        """List queryset should prefetch details when the details flag is enabled."""
        django_request = self.factory.get(self.url, {"entity_id": self.entity.id, "details": 1})
        django_request.entity = self.entity
        django_request.unit = "l"

        view = OperationViewSet()
        view.request = Request(django_request)
        view.action = "list"

        queryset = view.get_queryset()

        self.assertIn("details", queryset._prefetch_related_lookups)

    def test_correct_queryset_prefetches_details(self):
        """The correct action should prefetch details for the correction serializer."""
        django_request = self.factory.get(self.url, {"entity_id": self.entity.id})
        django_request.entity = self.entity
        django_request.unit = "l"

        view = OperationViewSet()
        view.request = Request(django_request)
        view.action = "correct"

        queryset = view.get_queryset()

        self.assertIn("details", queryset._prefetch_related_lookups)

    def test_export_queryset_prefetches_details(self):
        """The export action should prefetch details for the Excel export."""
        django_request = self.factory.get(self.url, {"entity_id": self.entity.id})
        django_request.entity = self.entity
        django_request.unit = "l"

        view = OperationViewSet()
        view.request = Request(django_request)
        view.action = "export_operations_to_excel"

        queryset = view.get_queryset()

        self.assertIn("details", queryset._prefetch_related_lookups)

    def test_export_operation_details_to_excel(self):
        self.operation_incorporation.durability_period = "202601"
        self.operation_incorporation.save(update_fields=["durability_period"])
        self.lot2.country_of_origin = None
        self.lot2.save(update_fields=["country_of_origin"])

        OperationDetail.objects.create(
            operation=self.operation_mac_bio,
            lot=self.lot1,
            volume=250,
            emission_rate_per_mj=15.0,
        )

        url = reverse(
            "operations-export-operation-details-to-excel",
            kwargs={"pk": self.operation_incorporation.id},
        )
        response = self.client.get(url, {"entity_id": self.entity.id})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.ms-excel")
        self.assertRegex(
            response["Content-Disposition"],
            rf'attachment; filename="tiruert_operation_{self.operation_incorporation.id}_details_\d{{8}}_\d{{6}}.xlsx"',
        )

        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        sheet = workbook["Détails opération"]
        rows = list(sheet.iter_rows(values_only=True))

        self.assertEqual(
            rows[0],
            (
                "ID Carbure",
                "Volume prélevé (L)",
                "Taux d'émission (gCO₂/MJ)",
                "Émissions évitées (tCO₂)",
                "Biocarburant",
                "Matière première",
                "Catégorie",
            ),
        )
        self.assertEqual(len(rows), 3)
        exported_rows = {row[0]: row for row in rows[1:]}
        detail = self.operation_incorporation.details.get(lot=self.lot1)
        self.assertEqual(
            exported_rows[self.lot1.carbure_id],
            (
                self.lot1.carbure_id,
                500,
                10.5,
                truncate(detail.avoided_emissions),
                self.lot1.biofuel.name,
                self.lot1.feedstock.name,
                self.lot1.feedstock.category,
            ),
        )
        self.assertIn(self.lot2.carbure_id, exported_rows)

    def test_list_operations_with_details(self):
        """Test GET /operations/?details=1 includes operation details."""
        query = {"entity_id": self.entity.id, "details": 1}
        response = self.client.get(self.url, query)

        self.assertEqual(response.status_code, 200)
        results = response.json()["results"]
        self.assertGreater(len(results), 0)
        self.assertIn("details", results[0])
        self.assertEqual(len(results[0]["details"]), 2)

    def test_list_operations_excludes_draft_credits(self):
        """Test GET /operations/ excludes draft operations that are credits."""
        # Create a draft operation that is a credit
        Operation.objects.create(
            type=Operation.CESSION,
            status=Operation.DRAFT,
            biofuel=self.biofuel_eth,
            credited_entity=self.entity,  # entity is credited, so it's a credit
            debited_entity=Entity.objects.exclude(id=self.entity.id).first(),
            to_depot=self.depot,
            customs_category=MatierePremiere.CONV,
            renewable_energy_share=1.0,
        )

        query = {"entity_id": self.entity.id}
        response = self.client.get(self.url, query)

        self.assertEqual(response.status_code, 200)
        # Should still be 4 because draft credit is excluded
        self.assertEqual(response.json()["count"], 4)

    def test_partial_update_operation_updates_and_returns_operation(self):
        """Test PATCH /operations/:id/ updates operation."""
        operation = Operation.objects.create(
            type=Operation.CESSION,
            status=Operation.PENDING,
            biofuel=self.biofuel_eth,
            debited_entity=self.entity,
            credited_entity=Entity.objects.exclude(id=self.entity.id).first(),
            to_depot=self.depot,
            customs_category=MatierePremiere.CONV,
            renewable_energy_share=1.0,
        )

        url = reverse("operations-detail", kwargs={"pk": operation.id})
        new_depot = Depot.objects.exclude(id=self.depot.id).first()
        payload = {"to_depot": new_depot.id}

        response = self.client.patch(
            url, data=payload, content_type="application/json", QUERY_STRING=f"entity_id={self.entity.id}"
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["to_depot"]["id"], new_depot.id)
        self.assertEqual(data["to_depot"]["name"], new_depot.name)

    def test_partial_update_status_changes_operation_status(self):
        """Test PATCH /operations/:id/ can update status."""
        operation = Operation.objects.create(
            type=Operation.CESSION,
            status=Operation.DRAFT,
            biofuel=self.biofuel_eth,
            debited_entity=self.entity,
            credited_entity=Entity.objects.exclude(id=self.entity.id).first(),
            to_depot=self.depot,
            customs_category=MatierePremiere.CONV,
            renewable_energy_share=1.0,
        )

        url = reverse("operations-detail", kwargs={"pk": operation.id})
        payload = {"status": Operation.PENDING}

        response = self.client.patch(
            url, data=payload, content_type="application/json", QUERY_STRING=f"entity_id={self.entity.id}"
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], Operation.PENDING)

    def test_destroy_operation_with_valid_conditions(self):
        """Test DELETE /operations/:id/ deletes operation when conditions met."""
        # Create operation that can be deleted
        operation = Operation.objects.create(
            type=Operation.CESSION,
            status=Operation.PENDING,
            biofuel=self.biofuel_eth,
            debited_entity=self.entity,
            credited_entity=Entity.objects.exclude(id=self.entity.id).first(),
            to_depot=self.depot,
            customs_category=MatierePremiere.CONV,
            renewable_energy_share=1.0,
        )

        url = reverse("operations-detail", kwargs={"pk": operation.id})

        response = self.client.delete(url, QUERY_STRING=f"entity_id={self.entity.id}")

        self.assertEqual(response.status_code, 204)
        self.assertFalse(Operation.objects.filter(id=operation.id).exists())

    def test_destroy_operation_with_invalid_type_returns_403(self):
        """Test DELETE /operations/:id/ returns 403 for non-deletable types."""
        # Create operation with type that cannot be deleted
        operation = Operation.objects.create(
            type=Operation.INCORPORATION,
            status=Operation.PENDING,
            biofuel=self.biofuel_eth,
            credited_entity=self.entity,
            to_depot=self.depot,
            customs_category=MatierePremiere.CONV,
            renewable_energy_share=1.0,
        )

        url = reverse("operations-detail", kwargs={"pk": operation.id})

        response = self.client.delete(url, QUERY_STRING=f"entity_id={self.entity.id}")

        self.assertEqual(response.status_code, 403)
        self.assertTrue(Operation.objects.filter(id=operation.id).exists())

    def test_destroy_operation_with_invalid_status_returns_403(self):
        """Test DELETE /operations/:id/ returns 403 for non-deletable statuses."""
        # Create operation with status that cannot be deleted
        operation = Operation.objects.create(
            type=Operation.CESSION,
            status=Operation.ACCEPTED,
            biofuel=self.biofuel_eth,
            debited_entity=self.entity,
            credited_entity=Entity.objects.exclude(id=self.entity.id).first(),
            to_depot=self.depot,
            customs_category=MatierePremiere.CONV,
            renewable_energy_share=1.0,
        )

        url = reverse("operations-detail", kwargs={"pk": operation.id})

        response = self.client.delete(url, QUERY_STRING=f"entity_id={self.entity.id}")

        self.assertEqual(response.status_code, 403)
        self.assertTrue(Operation.objects.filter(id=operation.id).exists())


class OperationViewSetDGECIntegrationTest(TestCase):
    """Integration tests for DGEC admin using selected_entity_id."""

    fixtures = [
        "json/biofuels.json",
        "json/feedstock.json",
        "json/countries.json",
        "json/entities.json",
        "json/depots.json",
        "json/entities_sites.json",
    ]

    @classmethod
    def setUpTestData(cls):
        cls.dgec_entity = Entity.objects.create(name="MTE - DGEC", entity_type=Entity.ADMIN)
        cls.operator = Entity.objects.filter(entity_type=Entity.OPERATOR).first()
        cls.other_operator = Entity.objects.filter(entity_type=Entity.OPERATOR).exclude(id=cls.operator.id).first()
        cls.depot = Depot.objects.first()
        cls.biofuel_eth = Biocarburant.objects.get(code="ETH")

        # Operations belonging to operator
        cls.op_credited = Operation.objects.create(
            type=Operation.INCORPORATION,
            status=Operation.VALIDATED,
            customs_category="CONV",
            biofuel=cls.biofuel_eth,
            credited_entity=cls.operator,
            to_depot=cls.depot,
            renewable_energy_share=1.0,
        )
        cls.op_debited = Operation.objects.create(
            type=Operation.CESSION,
            status=Operation.PENDING,
            customs_category="CONV",
            biofuel=cls.biofuel_eth,
            debited_entity=cls.operator,
            credited_entity=cls.other_operator,
            to_depot=cls.depot,
            renewable_energy_share=1.0,
        )
        # Operation unrelated to operator
        cls.op_other = Operation.objects.create(
            type=Operation.INCORPORATION,
            status=Operation.VALIDATED,
            customs_category="CONV",
            biofuel=cls.biofuel_eth,
            credited_entity=cls.other_operator,
            to_depot=cls.depot,
            renewable_energy_share=1.0,
        )

    def setUp(self):
        self.user = setup_current_user(
            self, "dgec@carbure.local", "DGEC", "password", [(self.dgec_entity, "ADMIN")], is_staff=True
        )
        self.url = reverse("operations-list")

    def test_dgec_admin_can_list_operations_for_target_entity(self):
        """DGEC admin using selected_entity_id should see only the target entity's operations."""
        response = self.client.get(
            self.url,
            {"entity_id": self.dgec_entity.id, "selected_entity_id": self.operator.id},
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()
        # Should see op_credited and op_debited, but not op_other
        self.assertEqual(data["count"], 2)

    def test_operator_cannot_use_selected_entity_id(self):
        """A regular operator cannot use selected_entity_id — should get 403."""
        # Re-login as operator
        setup_current_user(self, "op@carbure.local", "Op", "password", [(self.operator, "ADMIN")])

        response = self.client.get(
            self.url,
            {"entity_id": self.operator.id, "selected_entity_id": self.other_operator.id},
        )

        self.assertEqual(response.status_code, 403)
