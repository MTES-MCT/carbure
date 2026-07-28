from io import BytesIO
from unittest.mock import patch

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Entity
from core.tests_utils import setup_current_user
from entity.factories import EntityFactory
from tiruert.services.operation_excel_template import ENTITIES_SHEET_NAME, MAIN_SHEET_NAME, TABLE_HEADERS


class OperationExcelTemplateViewSetIntegrationTest(TestCase):
    """Integration tests for the operations import template endpoint."""

    def setUp(self):
        super().setUp()
        self.entity = EntityFactory.create(
            name="Template Operator",
            entity_type=Entity.OPERATOR,
            is_enabled=True,
            is_tiruert_liable=True,
            accise_number="ACC-100",
        )
        self.user = setup_current_user(self, "tester@carbure.local", "Tester", "gogogo", [(self.entity, "ADMIN")])
        self.url = reverse("operations-download-import-template")

    @patch("tiruert.services.operation_excel_template._get_operator_lots", return_value=[])
    @patch("tiruert.services.operation_excel_template.get_tiruert_operator_queryset", return_value=[])
    def test_download_import_template_wires_viewset_and_returns_excel_response(self, mock_get_entities, mock_get_lots):
        """Should call the template generator and return a downloadable Excel file."""
        response = self.client.get(self.url, QUERY_STRING=f"entity_id={self.entity.id}")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.ms-excel")
        self.assertIn('attachment; filename="tiruert_operations_import_template.xlsx"', response["Content-Disposition"])

        workbook = load_workbook(filename=BytesIO(response.content))
        self.addCleanup(workbook.close)
        main_sheet = workbook[MAIN_SHEET_NAME]
        entities_sheet = workbook[ENTITIES_SHEET_NAME]

        self.assertEqual(workbook.sheetnames, [MAIN_SHEET_NAME, ENTITIES_SHEET_NAME])
        self.assertEqual([cell.value for cell in main_sheet[1]], [label for label, _ in TABLE_HEADERS])
        self.assertEqual([cell.value for cell in main_sheet[2]], [key for _, key in TABLE_HEADERS])
        self.assertTrue(main_sheet.row_dimensions[2].hidden)
        self.assertEqual(entities_sheet[1][0].value, "Destinataire")
        self.assertEqual(entities_sheet[1][1].value, "Id")
        self.assertEqual(entities_sheet.max_row, 1)
        self.assertEqual(len(main_sheet.data_validations.dataValidation), 3)
